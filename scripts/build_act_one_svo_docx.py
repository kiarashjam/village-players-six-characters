#!/usr/bin/env python3
"""Build the ACT ONE actioning (SVO) worksheet as an editable Word file.

Act One of the Actor Rehearsal Script (same trims/strips/part headings as the
actor script), with an "Action: I ______ you" line under every speech so each
actor can write the subject-verb-object action for their lines — the
"I ___ you" work from the rehearsal map.

Output: outputs/act_one_svo.docx
Run: python scripts/build_act_one_svo_docx.py
"""
import os
from pathlib import Path
from bs4 import BeautifulSoup, NavigableString, Tag

import build_actor_script

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT_DIR = Path(os.environ.get("OUT_DIR", ROOT / "outputs"))
OUT_DIR.mkdir(parents=True, exist_ok=True)

GRAY = RGBColor(0x88, 0x80, 0x70)


def _runs_from(node, para, italic=False, bold=False):
    if isinstance(node, NavigableString):
        if str(node):
            r = para.add_run(str(node))
            r.italic = italic
            r.bold = bold
        return
    if isinstance(node, Tag):
        it = italic or node.name in ("em", "i")
        bd = bold or node.name in ("strong", "b")
        if node.name == "br":
            para.add_run().add_break()
            return
        for child in node.children:
            _runs_from(child, para, italic=it, bold=bd)


def build():
    soup = BeautifulSoup(build_actor_script.build_actor_html(), "html.parser")
    for sel in ("#act-2", "#act-3"):
        el = soup.select_one(sel)
        if el:
            el.decompose()

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)

    doc.add_paragraph("ACTIONING WORKSHEET — SUBJECT · VERB · OBJECT").alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Six Characters in Search of an Author — Act One")
    r.bold = True
    r.font.size = Pt(18)
    sub = doc.add_paragraph(
        "For every line: what are you DOING to the other person? Write one strong verb — "
        "“I ___ you” (I warn you, I coax you, I shame you). A verb aimed at someone, "
        "never a mood, never an adverb. One action per thought — long speeches turn; mark each turn.")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for rr in sub.runs:
        rr.italic = True
    doc.add_page_break()

    main = soup.select_one("main") or soup.body or soup
    for el in main.descendants:
        if not isinstance(el, Tag):
            continue
        if el.name == "section" and "act-header" in (el.get("class") or []):
            label = el.select_one(".label")
            doc.add_heading(label.get_text(" ", strip=True) if label else "Act", level=1)
            continue
        if el.name == "div" and "actor-part" in (el.get("class") or []):
            eyebrow = el.select_one(".part-eyebrow")
            title = el.select_one(".part-title")
            label = (eyebrow.get_text(" ", strip=True) if eyebrow else "")
            if title:
                label += (" — " if label else "") + title.get_text(" ", strip=True)
            doc.add_heading(label, level=2)
            continue
        if el.name == "p" and "speech" in (el.get("class") or []):
            para = doc.add_paragraph()
            sp = el.select_one(".speaker")
            speaker = sp.get_text(" ", strip=True) if sp else ""
            if speaker:
                rr = para.add_run(speaker.upper() + ".  ")
                rr.bold = True
            after = False
            for child in el.children:
                if isinstance(child, Tag) and "speaker" in (child.get("class") or []):
                    after = True
                    continue
                if not after and sp is not None:
                    continue
                if isinstance(child, Tag) and "action" in (child.get("class") or []):
                    r = para.add_run(child.get_text(" ", strip=True) + " ")
                    r.italic = True
                else:
                    _runs_from(child, para)
            # the SVO line
            act = doc.add_paragraph()
            act.paragraph_format.left_indent = Pt(36)
            act.paragraph_format.space_after = Pt(10)
            r1 = act.add_run("Action:  I ")
            r1.italic = True
            r1.font.color.rgb = GRAY
            r2 = act.add_run("_______________________")
            r2.font.color.rgb = GRAY
            r3 = act.add_run(" you")
            r3.italic = True
            r3.font.color.rgb = GRAY
            continue
        if el.name == "p" and "stage" in (el.get("class") or []):
            para = doc.add_paragraph()
            _runs_from(el, para)
            for rr in para.runs:
                rr.italic = True
            continue
        if el.name == "div" and "song" in (el.get("class") or []):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(24)
            _runs_from(el, para)
            for rr in para.runs:
                rr.italic = True
            continue
        if el.name == "div" and "curtain" in (el.get("class") or []):
            p = doc.add_paragraph(el.get_text(" ", strip=True))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for rr in p.runs:
                rr.bold = True

    out = OUT_DIR / "act_one_svo.docx"
    doc.save(str(out))
    speeches = len(soup.select("p.speech"))
    print(f"Done: {out} ({out.stat().st_size:,} bytes, {speeches} speeches with SVO lines)")


if __name__ == "__main__":
    build()
