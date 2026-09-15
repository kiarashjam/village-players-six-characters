#!/usr/bin/env python3
"""Build the Actor Rehearsal Script as a designed, editable Microsoft Word file.

Reuses build_actor_script.build_actor_html() (same trimming/stripping as the
PDF actor script), then walks the cleaned HTML and lays it out as a proper
Word document rather than a dump of plain paragraphs:

  - A4 page setup, generous margins, a footer with the page number.
  - Real Word styles (Normal, Heading 1/2/3, Speech, Stage Direction, Song,
    Curtain, Cast…) so the whole script can be restyled from Word's style
    pane, and the acts and parts show up in the Navigation pane.
  - A title page, a contents page, the cast, and the production note.
  - Speeches set with a hanging indent: the speaker in small caps and the
    house maroon, the stage business in italic, the line in plain serif.
  - Act openings with the roman numeral and the act's theme line.

Output: outputs/actor_script.docx
Run: python scripts/build_actor_script_docx.py
"""
import os
import re
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag

import build_actor_script  # same content pipeline as the PDF actor script

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT_DIR = Path(os.environ.get("OUT_DIR", ROOT / "outputs"))
OUT_DIR.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------- house style
BODY_FONT = "Georgia"          # warm serif, present on both Windows and macOS
INK = RGBColor(0x2A, 0x20, 0x1A)
MAROON = RGBColor(0x8B, 0x3A, 0x3A)
SOFT = RGBColor(0x6B, 0x5B, 0x48)
RULE = "C9BCA4"

SPEECH_INDENT = Cm(2.5)        # hanging indent for dialogue

# Order of children inside <w:pPr>, used to insert elements python-docx
# does not model (paragraph borders) at a schema-legal position.
_PPR_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


# ------------------------------------------------------------- xml utilities
def _rpr(obj):
    """Return the <w:rPr> of a run or a style, creating it if needed."""
    el = obj._element if hasattr(obj, "_element") else obj._r
    return el.get_or_add_rPr()


def _set_font(obj, name=BODY_FONT):
    """Pin a run's/style's font, clearing any theme font that would win."""
    rpr = _rpr(obj)
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(theme)) is not None:
            del rf.attrib[qn(theme)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), name)


def _flag(obj, tag, on=True):
    """Set a simple boolean run property such as <w:smallCaps/>."""
    rpr = _rpr(obj)
    el = rpr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        rpr.append(el)
    el.set(qn("w:val"), "1" if on else "0")


def _tracking(obj, twentieths):
    """Letter-spacing, in twentieths of a point (30 = 1.5pt)."""
    rpr = _rpr(obj)
    el = rpr.find(qn("w:spacing"))
    if el is None:
        el = OxmlElement("w:spacing")
        rpr.append(el)
    el.set(qn("w:val"), str(int(twentieths)))


def _color(obj, rgb):
    rpr = _rpr(obj)
    el = rpr.find(qn("w:color"))
    if el is None:
        el = OxmlElement("w:color")
        rpr.append(el)
    el.set(qn("w:val"), str(rgb))
    if el.get(qn("w:themeColor")) is not None:
        del el.attrib[qn("w:themeColor")]


def _bottom_rule(para, color=RULE, size=4, space=8):
    """A hairline under a paragraph (used beneath part titles)."""
    ppr = para._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.insert_element_before(pbdr, *_PPR_AFTER_PBDR)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))        # eighths of a point
    bottom.set(qn("w:space"), str(space))    # points between text and rule
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def _field(para, instruction):
    """Insert a Word field (used for PAGE / NUMPAGES in the footer)."""
    r1 = para.add_run()._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    r1.append(begin)
    r2 = para.add_run()._r
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    r2.append(instr)
    r3 = para.add_run()._r
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r3.append(end)


# ------------------------------------------------------------------- styles
def _new_style(doc, name, base="Normal"):
    from docx.enum.style import WD_STYLE_TYPE
    if name in [s.name for s in doc.styles]:
        return doc.styles[name]
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles[base]
    return st


def build_styles(doc):
    """Define every style the script uses, so Word can restyle it wholesale."""
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    _set_font(normal)
    pf = normal.paragraph_format
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.18
    pf.widow_control = True

    # --- Act opening (Heading 1: shows in Word's Navigation pane)
    h1 = doc.styles["Heading 1"]
    h1.font.size = Pt(13)
    h1.font.bold = False
    h1.font.italic = False
    h1.font.color.rgb = MAROON
    _set_font(h1)
    _flag(h1, "w:smallCaps")
    _tracking(h1, 60)
    p = h1.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(0)
    p.space_after = Pt(4)
    p.keep_with_next = True

    # --- Part title (Heading 2)
    h2 = doc.styles["Heading 2"]
    h2.font.size = Pt(15)
    h2.font.bold = False
    h2.font.italic = False
    h2.font.color.rgb = MAROON
    _set_font(h2)
    p = h2.paragraph_format
    p.space_before = Pt(2)
    p.space_after = Pt(12)
    p.keep_with_next = True

    # --- Section heads inside the front matter (Heading 3)
    h3 = doc.styles["Heading 3"]
    h3.font.size = Pt(12)
    h3.font.bold = False
    h3.font.color.rgb = MAROON
    _set_font(h3)
    _flag(h3, "w:smallCaps")
    _tracking(h3, 30)
    p = h3.paragraph_format
    p.space_before = Pt(16)
    p.space_after = Pt(6)
    p.keep_with_next = True

    # --- Dialogue: hanging indent under the speaker name
    sp = _new_style(doc, "Speech")
    p = sp.paragraph_format
    p.left_indent = SPEECH_INDENT
    p.first_line_indent = -SPEECH_INDENT
    p.space_after = Pt(8)
    p.keep_together = False
    p.widow_control = True

    # --- Stage directions on their own line
    sd = _new_style(doc, "Stage Direction")
    sd.font.size = Pt(10.5)
    sd.font.italic = True
    sd.font.color.rgb = SOFT
    p = sd.paragraph_format
    p.left_indent = SPEECH_INDENT
    p.right_indent = Cm(1.2)
    p.space_before = Pt(8)
    p.space_after = Pt(10)
    p.keep_together = True

    # --- The song in Act Two
    sg = _new_style(doc, "Song")
    sg.font.size = Pt(10.5)
    sg.font.italic = True
    sg.font.color.rgb = SOFT
    p = sg.paragraph_format
    p.left_indent = Cm(3.6)
    p.space_before = Pt(10)
    p.space_after = Pt(12)
    p.keep_together = True

    # --- "— End of Act I —"
    ct = _new_style(doc, "Curtain")
    ct.font.size = Pt(11)
    ct.font.color.rgb = MAROON
    _flag(ct, "w:smallCaps")
    _tracking(ct, 60)
    p = ct.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(22)
    p.space_after = Pt(4)
    p.keep_together = True

    # --- Small caps eyebrow above a part title ("Part I of Act One")
    eb = _new_style(doc, "Part Eyebrow")
    eb.font.size = Pt(9)
    eb.font.color.rgb = SOFT
    _flag(eb, "w:smallCaps")
    _tracking(eb, 60)
    p = eb.paragraph_format
    p.space_before = Pt(26)
    p.space_after = Pt(0)
    p.keep_with_next = True

    # --- Centred italic lines (act theme, cover credits, contents intro)
    cen = _new_style(doc, "Centred Note")
    cen.font.italic = True
    cen.font.color.rgb = SOFT
    p = cen.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(6)
    p.space_after = Pt(6)

    # --- Cast and contents lines
    cl = _new_style(doc, "Cast Line")
    p = cl.paragraph_format
    p.left_indent = Cm(0.8)
    p.first_line_indent = Cm(-0.4)
    p.space_after = Pt(5)

    tc = _new_style(doc, "Contents Line")
    p = tc.paragraph_format
    p.left_indent = Cm(1.6)
    p.first_line_indent = Cm(-1.6)
    p.space_after = Pt(4)

    # --- Body paragraphs in the production note
    nb = _new_style(doc, "Note Body")
    nb.font.size = Pt(10.5)
    p = nb.paragraph_format
    p.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc


def page_setup(doc):
    sec = doc.sections[0]
    sec.start_type = WD_SECTION_START.NEW_PAGE
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.2)
    sec.footer_distance = Cm(1.2)
    sec.different_first_page_header_footer = True  # no folio on the title page

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Six Characters in Search of an Author  ·  Actor Rehearsal Script  ·  ")
    _field(footer, " PAGE ")
    for r in footer.runs:
        r.font.size = Pt(8.5)
        r.font.color.rgb = SOFT
        _set_font(r)
    _tracking(run, 15)
    return doc


# ------------------------------------------------------------- pre-processing
def upright_brackets(soup):
    """Split the [ ] off the stage business so the brackets can stay upright.

    Published play texts set the brackets roman and only the words inside in
    italic; an italic closing bracket leans into the word after it. Each
    <span class="action">[…]</span> becomes
    <span class="action"><span class="bracket">[</span>…<span class="bracket">]</span></span>.
    """
    n = 0
    for sp in soup.select("span.action"):
        strings = [x for x in sp.descendants if isinstance(x, NavigableString)]
        if not strings:
            continue
        first, last = strings[0], strings[-1]
        if not str(first).lstrip().startswith("[") or not str(last).rstrip().endswith("]"):
            continue
        if first is last:
            inner = str(first).strip()
            first.replace_with(NavigableString(inner[1:-1]))
        else:
            first.replace_with(NavigableString(str(first).lstrip()[1:]))
            last.replace_with(NavigableString(str(last).rstrip()[:-1]))
        for pos, char in ((0, "["), (None, "]")):
            b = soup.new_tag("span", **{"class": "bracket"})
            b.string = char
            if pos == 0:
                sp.insert(0, b)
            else:
                sp.append(b)
        n += 1
    return n


# --------------------------------------------------------------- text runs
class Writer:
    """Emits runs into a paragraph, tidying the whitespace as it goes.

    The source markup puts the sentence's opening full stop after the stage
    business — "<span class=speaker>The Manager</span> <span class=action>[…]</span>. No."
    — so runs have to be joined without leaving a space in front of it.
    """

    def __init__(self, para):
        self.para = para

    def ensure_space(self):
        """A single separating space, unless the line already ends with one."""
        runs = self.para.runs
        if runs and runs[-1].text and not runs[-1].text.endswith(" "):
            r = self.para.add_run(" ")
            _set_font(r)

    def add(self, text, italic=False, bold=False, color=None):
        text = re.sub(r"\s+", " ", text)
        if not text:
            return
        runs = self.para.runs
        prev = runs[-1].text if runs else ""
        if text.startswith(" ") and (not prev or prev[-1].isspace()):
            text = text.lstrip()
        if not text:
            return
        if text[0] in ".,;:!?…)]" and prev.endswith(" "):
            runs[-1].text = prev.rstrip()
        r = self.para.add_run(text)
        if italic:
            r.italic = True
        if bold:
            r.bold = True
        _set_font(r)
        if color is not None:
            r.font.color.rgb = color
        return r

    def walk(self, node, italic=False, bold=False, color=None):
        """Render an inline node, honouring <em>/<i>/<strong>/<b>/<br>."""
        if isinstance(node, NavigableString):
            self.add(str(node), italic=italic, bold=bold, color=color)
            return
        if not isinstance(node, Tag):
            return
        if node.name == "br":
            if self.para.runs:
                self.para.runs[-1].add_break()
            return
        classes = node.get("class") or []
        it = italic or node.name in ("em", "i") or "action" in classes
        bd = bold or node.name in ("strong", "b")
        col = SOFT if "action" in classes else color
        if "bracket" in classes:
            it = False
        if "action" in classes:
            self.ensure_space()
        for child in node.children:
            self.walk(child, italic=it, bold=bd, color=col)


def _text(el, default=""):
    return el.get_text(" ", strip=True) if el is not None else default


# ------------------------------------------------------------- front matter
def title_page(doc, soup):
    cover = soup.select_one("section.cover")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(104)

    eyebrow = doc.add_paragraph(style="Part Eyebrow")
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eyebrow.paragraph_format.space_before = Pt(0)
    r = eyebrow.add_run(_text(cover.select_one(".eyebrow"), "Actor Rehearsal Script"))
    _set_font(r)
    r.font.size = Pt(11)
    r.font.color.rgb = MAROON
    _flag(r, "w:smallCaps")
    _tracking(r, 90)

    for line, size in (("Six Characters", 34), ("in Search of an Author", 34)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8 if size == 34 else 0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        _set_font(r)
        r.font.size = Pt(size)
        r.font.color.rgb = MAROON

    ital = doc.add_paragraph(style="Centred Note")
    ital.paragraph_format.space_before = Pt(10)
    r = ital.add_run(_text(cover.select_one(".cover-italian"), "Sei personaggi in cerca d'autore"))
    _set_font(r)
    r.font.size = Pt(12)

    orn = doc.add_paragraph(style="Centred Note")
    orn.paragraph_format.space_before = Pt(22)
    r = orn.add_run("·  ·  ·")
    _set_font(r)
    r.italic = False
    r.font.size = Pt(14)
    r.font.color.rgb = MAROON

    sub = doc.add_paragraph(style="Part Eyebrow")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_before = Pt(18)
    r = sub.add_run(_text(cover.select_one(".subtitle"), "A Comedy in the Making"))
    _set_font(r)
    r.font.size = Pt(11)
    r.font.color.rgb = INK
    _flag(r, "w:smallCaps")
    _tracking(r, 70)

    credit = doc.add_paragraph(style="Centred Note")
    credit.paragraph_format.space_before = Pt(46)
    credit.paragraph_format.space_after = Pt(0)
    credit.paragraph_format.line_spacing = 1.5
    for i, line in enumerate([
            "Luigi Pirandello  ·  trans. Edward Storer (1922)",
            "Rewritten and directed by Kiarash Jamshidi",
            "Village Players, Lausanne"]):
        r = credit.add_run(line)
        _set_font(r)
        r.font.size = Pt(11)
        r.italic = True
        r.font.color.rgb = SOFT if i != 1 else INK
        if i < 2:
            r.add_break()

    foot = doc.add_paragraph(style="Centred Note")
    foot.paragraph_format.space_before = Pt(72)
    r = foot.add_run("Working script for rehearsal. Dialogue, stage directions, and cues only.")
    _set_font(r)
    r.font.size = Pt(9.5)
    doc.add_page_break()


def contents_page(doc, soup):
    h = doc.add_paragraph(style="Heading 3")
    h.paragraph_format.space_before = Pt(0)
    h.add_run("Contents")

    body = soup.select_one("main") or soup
    current_act = None
    for el in body.descendants:
        if not isinstance(el, Tag):
            continue
        classes = el.get("class") or []
        if el.name == "section" and "act-header" in classes:
            current_act = _text(el.select_one(".label"), "Act")
            p = doc.add_paragraph(style="Contents Line")
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(4.6))
            r = p.add_run(current_act)
            _set_font(r)
            r.font.color.rgb = MAROON
            _flag(r, "w:smallCaps")
            _tracking(r, 40)
            theme = _text(el.select_one(".act-theme"))
            if theme:
                r2 = p.add_run("\t" + theme)
                _set_font(r2)
                r2.italic = True
                r2.font.size = Pt(9.5)
                r2.font.color.rgb = SOFT
        elif el.name == "div" and "actor-part" in classes:
            eyebrow = _text(el.select_one(".part-eyebrow"))
            title = _text(el.select_one(".part-title"))
            num = eyebrow.split(" of ")[0].replace("Part ", "").strip() if eyebrow else ""
            p = doc.add_paragraph(style="Contents Line")
            pf = p.paragraph_format
            pf.left_indent = Cm(3.0)
            pf.first_line_indent = Cm(-1.4)
            pf.tab_stops.add_tab_stop(Cm(3.0))
            r = p.add_run(f"{num}\t" if num else "\t")
            _set_font(r)
            r.font.color.rgb = SOFT
            r2 = p.add_run(title)
            _set_font(r2)
    doc.add_page_break()


def cast_pages(doc, soup):
    cast = soup.select_one("section.cast")
    if cast is None:
        return
    h = doc.add_paragraph(style="Heading 3")
    h.paragraph_format.space_before = Pt(0)
    h.add_run(_text(cast.select_one("h2"), "Dramatis Personæ"))
    sub = _text(cast.select_one("p.cast-sub"))
    if sub:
        p = doc.add_paragraph(style="Centred Note")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(sub)
        _set_font(r)
        r.font.size = Pt(10)

    for group in cast.select("div.cast-group"):
        gh = group.select_one(".cast-h3")
        if gh is not None:
            p = doc.add_paragraph(style="Heading 3")
            p.add_run(_text(gh))
        for li in group.select("li"):
            para = doc.add_paragraph(style="Cast Line")
            extra = li.select_one(".silent") or li.select_one(".role-list")
            name = li.get_text(" ", strip=True)
            if extra is not None:
                name = name.replace(_text(extra), "").strip(" —·")
            r = para.add_run(name)
            _set_font(r)
            r.bold = True
            if extra is not None:
                r2 = para.add_run("   " + _text(extra))
                _set_font(r2)
                r2.italic = True
                r2.font.size = Pt(9.5)
                r2.font.color.rgb = SOFT
    doc.add_page_break()


def production_note(doc, soup):
    prod = soup.select_one("section.actor-production-note")
    if prod is None:
        return
    h = doc.add_paragraph(style="Heading 3")
    h.paragraph_format.space_before = Pt(0)
    h.add_run(_text(prod.select_one("h2"), "Production note"))
    for p in prod.select("p"):
        para = doc.add_paragraph(style="Note Body")
        Writer(para).walk(p)
    doc.add_page_break()


# ------------------------------------------------------------------- the play
def act_opening(doc, el, first):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.page_break_before = not first
    spacer.paragraph_format.space_after = Pt(72)

    doc.add_paragraph(_text(el.select_one(".label"), "Act"), style="Heading 1")

    numeral = _text(el.select_one(".numeral"))
    if numeral:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(numeral)
        _set_font(r)
        r.font.size = Pt(46)
        r.font.color.rgb = MAROON

    p = doc.add_paragraph(style="Centred Note")
    r = p.add_run("·  ·  ·")
    r.italic = False
    _set_font(r)
    r.font.size = Pt(14)
    r.font.color.rgb = MAROON

    theme = _text(el.select_one(".act-theme"))
    if theme:
        p = doc.add_paragraph(style="Centred Note")
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(30)
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.right_indent = Cm(1.5)
        r = p.add_run(theme)
        _set_font(r)
        r.font.size = Pt(11.5)


def part_heading(doc, el):
    eyebrow = _text(el.select_one(".part-eyebrow"))
    title = _text(el.select_one(".part-title"))
    if eyebrow:
        p = doc.add_paragraph(style="Part Eyebrow")
        p.add_run(eyebrow)
    head = doc.add_paragraph(title, style="Heading 2")
    _bottom_rule(head)


def speech(doc, el):
    para = doc.add_paragraph(style="Speech")
    w = Writer(para)
    sp = el.select_one(".speaker")
    if sp is not None:
        r = para.add_run(_text(sp))
        _set_font(r)
        r.bold = True
        r.font.color.rgb = MAROON
        _flag(r, "w:smallCaps")
        _tracking(r, 10)
    seen_speaker = sp is None
    first = True
    for child in el.children:
        if not seen_speaker:
            if isinstance(child, Tag) and "speaker" in (child.get("class") or []):
                seen_speaker = True
            continue
        # "The Manager. Save your wit…" — keep the stop tight against the name,
        # so the name's letter-spacing falls after it rather than before.
        if first and isinstance(child, NavigableString):
            m = re.match(r"^([.,;:!?\u2026]+)", str(child))
            if m:
                r = para.add_run(m.group(1))
                _set_font(r)
                r.bold = True
                r.font.color.rgb = MAROON
                child = NavigableString(str(child)[m.end():])
        first = False
        w.walk(child)


def build_body(doc, soup):
    main = soup.select_one("main") or soup.body or soup
    first_act = True
    for el in main.descendants:
        if not isinstance(el, Tag):
            continue
        classes = el.get("class") or []
        if el.name == "section" and "act-header" in classes:
            act_opening(doc, el, first_act)
            first_act = False
        elif el.name == "div" and "actor-part" in classes:
            part_heading(doc, el)
        elif el.name == "p" and "speech" in classes:
            speech(doc, el)
        elif el.name == "p" and "stage" in classes:
            para = doc.add_paragraph(style="Stage Direction")
            Writer(para).walk(el)
        elif el.name == "div" and "song" in classes:
            para = doc.add_paragraph(style="Song")
            Writer(para).walk(el)
        elif el.name == "div" and "curtain" in classes:
            doc.add_paragraph(_text(el), style="Curtain")


def build():
    html = build_actor_script.build_actor_html()
    soup = BeautifulSoup(html, "html.parser")
    upright_brackets(soup)

    doc = Document()
    build_styles(doc)
    page_setup(doc)

    props = doc.core_properties
    props.title = "Six Characters in Search of an Author — Actor Rehearsal Script"
    props.author = "Kiarash Jamshidi"
    props.subject = "Village Players, Lausanne"
    props.comments = "Rewritten and directed by Kiarash Jamshidi, after Pirandello (trans. Edward Storer, 1922)."

    title_page(doc, soup)
    contents_page(doc, soup)
    cast_pages(doc, soup)
    production_note(doc, soup)
    build_body(doc, soup)

    out = OUT_DIR / "actor_script.docx"
    doc.save(str(out))
    speeches = len(soup.select("p.speech"))
    parts = len(soup.select("div.actor-part"))
    print(f"Done: {out} ({out.stat().st_size:,} bytes, "
          f"{speeches} speeches, {parts} parts, {len(doc.paragraphs)} paragraphs)")


if __name__ == "__main__":
    build()
